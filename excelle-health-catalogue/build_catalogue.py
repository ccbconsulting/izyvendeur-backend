# -*- coding: utf-8 -*-
# Construit le fichier de reference "Catalogue Excelle Health by Vestige" a partir des grilles de prix
# Vestige fournies (photos), classees par article (une ligne = un article distinct - aucune vraie variante
# couleur/taille dans ces donnees, la taille/le format est deja dans le nom du produit). Prix retenu : SP
# ("prix de vente" / prix consommateur final), jamais PV (point de valeur MLM, pas un prix) ni DP (prix
# distributeur/depot). Ce fichier est une reference de saisie manuelle dans /admin > Catalogue (IzyVendeur
# n'a pas d'import en masse) - pas un fichier destine a etre re-importe automatiquement.
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

wb = openpyxl.Workbook()

# ---------------------------------------------------------------------------------------------------
# Onglet 1 : instructions
# ---------------------------------------------------------------------------------------------------
ws0 = wb.active
ws0.title = "Instructions"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 100

titre = ws0.cell(row=1, column=1, value="Catalogue Excelle Health by Vestige — reference pour saisie dans /admin")
titre.font = Font(name=FONT_NAME, size=14, bold=True)

lignes = [
    "",
    "Ce fichier reprend les grilles de prix Vestige que vous avez transmises, classees par article",
    "(un article = une ligne du tableau ; aucun de ces produits n'a de vraie variante couleur/taille,",
    "le format/le poids est deja indique dans le nom).",
    "",
    "Prix retenu : la colonne \"SP\" (prix de vente / prix consommateur final), jamais la colonne \"PV\"",
    "(point de valeur, un indicateur de commissionnement Vestige, pas un prix) ni la colonne \"DP\" (prix",
    "distributeur/depot, plus bas, reserve aux distributeurs Vestige).",
    "",
    "IzyVendeur n'a pas encore d'import en masse du catalogue : ce fichier sert de reference pendant que",
    "vous creez chaque article a la main dans /admin > Catalogue > \"Ajouter un article\", avec :",
    "  - Nom = colonne \"Nom de l'article\"",
    "  - Categorie = colonne \"Categorie\" (regroupement propose, modifiable)",
    "  - Couleur / Taille = mettre \"-\" / \"Unique\" (champs obligatoires cote IzyVendeur, sans usage reel ici)",
    "  - Prix = colonne \"Prix de vente (FCFA)\"",
    "  - Stock reel = a completer vous-meme (colonne \"Stock initial\" laissee vide - la grille de prix Vestige",
    "    ne donne aucune quantite en stock, seulement les tarifs)",
    "",
    "Points a verifier avant saisie (signales aussi dans la colonne \"Remarque\") :",
    "  1) \"Vestige Prime Absorvit Vitamin C Sublingual Spray\" : une correction manuscrite apparait sur une",
    "     des pages a cote du prix imprime (18 000). Confirmez le prix exact avant de le saisir.",
    "  2) \"Fiber\" (page Ayurvedique, 21 000 FCFA) et \"Vestige Fibre 200 g Jar\" (17 000 FCFA) : noms proches",
    "     mais prix differents - verifiez s'il s'agit bien de deux produits distincts (format different ?)",
    "     avant de creer les deux.",
    "  3) \"Vestige Colostrum\" apparaissait deux fois dans vos documents (memes DP/SP) - comptabilise une",
    "     seule fois ici.",
    "",
    f"Total : 71 articles repartis en 8 categories.",
]
for i, texte in enumerate(lignes, start=2):
    c = ws0.cell(row=i, column=1, value=texte)
    c.font = Font(name=FONT_NAME, size=11, bold=texte.startswith("Points a verifier") or texte.startswith("Prix retenu") or texte.startswith("IzyVendeur"))
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------------------------------
# Onglet 2 : catalogue
# ---------------------------------------------------------------------------------------------------
ws = wb.create_sheet("Catalogue")

HEADERS = ["Categorie", "Nom de l'article", "Couleur (variante)", "Taille (variante)", "Prix de vente (FCFA)", "Stock initial", "Seuil d'alerte", "Remarque"]

# (categorie, nom, prix, remarque)
DONNEES = [
    # --- Vestige - Complements & bien-etre ---
    ("Vestige - Complements & bien-etre", "Vestige Colostrum", 17000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Prime Concentrated Mineral Drops (CMD)", 25000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Shatavari Max", 18500, ""),
    ("Vestige - Complements & bien-etre", "Vestige Cranberry 60 Capsules", 25000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Fibre 200 g Jar", 17000, "Verifier si distinct de \"Fiber\" (Ayurvedique, 21 000 FCFA)"),
    ("Vestige - Complements & bien-etre", "Vestige Prime Energy Booster", 23000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Collagen", 15000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Veg Collagen with Rosehip Extract & Vitamin C", 15000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Prime Absorvit Vitamin C Sublingual Spray", 18000, "Prix a confirmer - correction manuscrite vue sur une des pages"),
    ("Vestige - Complements & bien-etre", "Vestige Prime Multi Vita+Min Gummies 60N", 13000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Prime Absorvit Melatonin", 19000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Prime Absorvit Vitamin B12", 19000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Prime Absorvit Biotin Sublingual Spray 14ml", 13000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Curcumin Plus", 23000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Prime Sea Buckthorn", 23000, ""),
    ("Vestige - Complements & bien-etre", "Vestige Stevia", 4000, ""),

    # --- Ayusante ---
    ("Ayusante", "Ayusante Toxclean", 18000, ""),
    ("Ayusante", "Ayusante Kidneyhealth", 18000, ""),
    ("Ayusante", "Ayusante Vital Complex", 18000, ""),
    ("Ayusante", "Ayusante Respocare", 18000, ""),
    ("Ayusante", "Ayusante Procard 60 Capsules", 15000, ""),
    ("Ayusante", "Ayusante Prostate Care", 17500, ""),
    ("Ayusante", "Ayusante Gluco Health", 15000, ""),
    ("Ayusante", "Ayusante Liverhealth", 15000, ""),

    # --- Autres marques bien-etre ---
    ("Autres marques bien-etre", "Vescare Insta Relief Cream", 5000, ""),
    ("Autres marques bien-etre", "Invigo Health Drops Panch Tulsi Ark 30ml", 7500, ""),

    # --- Ayurvedique - Sante ---
    ("Ayurvedique - Sante", "NONI 500mg", 12500, ""),
    ("Ayurvedique - Sante", "Spirulina 500 mg", 11000, ""),
    ("Ayurvedique - Sante", "Amla 500 mg", 8000, ""),
    ("Ayurvedique - Sante", "Aloe Vera 500 mg", 8000, ""),
    ("Ayurvedique - Sante", "Neem 0.25 ml", 9000, ""),
    ("Ayurvedique - Sante", "Ganoderma 425 mg", 18500, ""),
    ("Ayurvedique - Sante", "Flax Oil 500mg", 12500, ""),
    ("Ayurvedique - Sante", "Glucosamine 450 mg", 14500, ""),
    ("Ayurvedique - Sante", "Folic & Iron Plus", 7500, ""),
    ("Ayurvedique - Sante", "Prime Krill Oil 500 mg 30", 22500, ""),
    ("Ayurvedique - Sante", "Calcium 625 mg 100", 7500, ""),
    ("Ayurvedique - Sante", "Coenzyme Q10", 17000, ""),
    ("Ayurvedique - Sante", "Veslim Shake - Mango Flavour", 24000, ""),
    ("Ayurvedique - Sante", "Grand Bary", 36000, ""),
    ("Ayurvedique - Sante", "Fiber", 21000, "Verifier si distinct de \"Vestige Fibre 200 g Jar\" (17 000 FCFA)"),
    ("Ayurvedique - Sante", "L-Argine 10g", 31000, ""),
    ("Ayurvedique - Sante", "Vestige Veslim Tea and Green Coffee", 22000, ""),
    ("Ayurvedique - Sante", "Vestige Veslim Capsule", 23000, ""),
    ("Ayurvedique - Sante", "Eye Support", 17000, ""),
    ("Ayurvedique - Sante", "Detox Foot Patches", 22500, ""),
    ("Ayurvedique - Sante", "Prime Combiotics", 16500, ""),

    # --- Agriculture ---
    ("Agriculture", "Engrais Agri 82 500ml", 21000, ""),
    ("Agriculture", "Engrais Agri Humic 500ml", 21000, ""),
    ("Agriculture", "Engrais Agri Humic Granules 5kg", 31000, ""),

    # --- Assure - Soins & beaute ---
    ("Assure - Soins & beaute", "Assure Complexion Bar 75g", 2000, ""),
    ("Assure - Soins & beaute", "Assure Daily Moisturiser 250ml", 8000, ""),
    ("Assure - Soins & beaute", "Assure Body Wash", 7000, ""),
    ("Assure - Soins & beaute", "Assure Hand Wash", 4000, ""),
    ("Assure - Soins & beaute", "Dentassure Gano Toothpaste 100g", 3500, ""),
    ("Assure - Soins & beaute", "Assure Purifying Cleanser + Toner 250ml", 8000, ""),
    ("Assure - Soins & beaute", "Parfum poche (Assure Blossom Perfume)", 3500, ""),
    ("Assure - Soins & beaute", "Assure Perfume Spray Artic", 8000, ""),
    ("Assure - Soins & beaute", "Assure Rapture Deo Women", 8000, ""),
    ("Assure - Soins & beaute", "Assure Mild Exfoliating Face Scrub 60g", 6500, ""),
    ("Assure - Soins & beaute", "Assure Anti-Ageing Night Cream 60g", 6500, ""),
    ("Assure - Soins & beaute", "Assure Clarifying Face Wash 60g", 5000, ""),
    ("Assure - Soins & beaute", "DewGarden Sanitary (Serviette Hygienique)", 3500, ""),
    ("Assure - Soins & beaute", "Poudre Femme (Echeant Body Talc)", 3000, ""),
    ("Assure - Soins & beaute", "Foot Cream", 5000, ""),
    ("Assure - Soins & beaute", "Assure Soap", 2000, ""),

    # --- Entretien ---
    ("Entretien", "Ultra Wash Liquid Laundry Detergent 500ml", 10000, ""),
    ("Entretien", "Ultra Guard Toilet Cleaner 500ml", 5000, ""),
    ("Entretien", "Ultra Scrub Dishwashing Liquid 500ml", 5000, ""),

    # --- Dietetique ---
    ("Dietetique", "Zeta Tea", 7000, ""),
    ("Dietetique", "Zeta Cofee", 5000, ""),
]

# En-tetes
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for col, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

cat_fills = {}
palette = ["DCE6F1", "FDE9D9", "E2EFDA", "FFF2CC", "F2DCDB", "EAD1DC", "D9E1F2", "E2F0D9"]
categories_uniques = []
for cat, *_ in DONNEES:
    if cat not in categories_uniques:
        categories_uniques.append(cat)
for i, cat in enumerate(categories_uniques):
    cat_fills[cat] = PatternFill(start_color=palette[i % len(palette)], end_color=palette[i % len(palette)], fill_type="solid")

row = 2
for cat, nom, prix, remarque in DONNEES:
    ws.cell(row=row, column=1, value=cat).font = Font(name=FONT_NAME, size=10)
    ws.cell(row=row, column=2, value=nom).font = Font(name=FONT_NAME, size=10)
    ws.cell(row=row, column=3, value="-").font = Font(name=FONT_NAME, size=10)
    ws.cell(row=row, column=4, value="Unique").font = Font(name=FONT_NAME, size=10)
    pc = ws.cell(row=row, column=5, value=prix)
    pc.font = Font(name=FONT_NAME, size=10)
    pc.number_format = "#,##0"
    ws.cell(row=row, column=6, value=None).font = Font(name=FONT_NAME, size=10)
    ws.cell(row=row, column=7, value=2).font = Font(name=FONT_NAME, size=10)
    rc = ws.cell(row=row, column=8, value=remarque)
    rc.font = Font(name=FONT_NAME, size=10, italic=True, color="C00000") if remarque else Font(name=FONT_NAME, size=10)

    fill = cat_fills[cat]
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        cell.border = border
        if col == 1:
            cell.fill = fill
    row += 1

total_row = row + 1
ws.cell(row=total_row, column=1, value="Total articles").font = Font(name=FONT_NAME, size=10, bold=True)
tc = ws.cell(row=total_row, column=2, value=f"=COUNTA(B2:B{row-1})")
tc.font = Font(name=FONT_NAME, size=10, bold=True)

widths = [30, 46, 16, 16, 18, 14, 14, 46]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:H{row-1}"

wb.save("Catalogue_ExcelleHealth_Vestige.xlsx")
print("OK - fichier ecrit")
